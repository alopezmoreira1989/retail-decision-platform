"""Tests for Phase 2 Vertical Slice #1 (POS + Assortment feeds,
Profile A/Profile B, against RETAILER-001's existing Phase 1
snapshot). Covers the 24 testable invariants from the approved design
brief and Implementation GO, provenance checks for both feeds, and the
Source & Delivery Model checkpoint's additions (source_family,
delivery_mechanism, column casing, retailer-native field names, and
the new operational-baggage columns).

All tests read the real, already-generated Phase 1 snapshot at
SNAPSHOT_DIR read-only; none of them regenerate or mutate it.
"""

from __future__ import annotations

import csv
import hashlib
from collections import defaultdict
from datetime import datetime, time, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from simulator.feeds.assortment import (
    ASSORTMENT_COLUMN_ORDER,
    generate_assortment_artifacts,
    weekly_delivery_periods,
)
from simulator.feeds.identifiers import resolve_store_identifier
from simulator.feeds.pos import POS_COLUMN_ORDER, generate_pos_artifacts
from simulator.feeds.profiles import (
    COLUMN_CASING_LOWER_SNAKE_CASE,
    DUPLICATE_POS_DELIVERY_OFFSET,
    EXCESS_LATENCY_ELIGIBLE_OFFSETS,
    EXCESS_LATENCY_EXTRA_DAYS,
    FEED_DELIVERY_MECHANISM,
    FEED_SOURCE_FAMILY,
    MISSING_ASSORTMENT_DELIVERY_OFFSET,
    MISSING_POS_DELIVERY_OFFSETS,
    STALE_ASSORTMENT_VISIBILITY_EXTRA_DAYS,
    FeedProfile,
    RetailerFeedProfile,
    build_identifier_scheme,
    build_profile_a,
    build_profile_b,
)
from simulator.feeds.runner import (
    generate_profile_feeds,
    run_vertical_slice_1,
    write_generated_feeds,
)
from simulator.feeds.snapshot_reader import _DAILY_STATE_COLUMNS, read_phase1_snapshot

SNAPSHOT_DIR = (
    Path(__file__).resolve().parents[3] / "simulator" / "world" / "output" / "dev_slice" / "seed42"
)
NEVER_CARRIED_PAIR = ("STORE-001", "SKU-003")  # matches world's own test fixture naming
TEMPORAL_GAP_PAIR = (
    "STORE-001",
    "SKU-002",
)  # the one real physical_assortment change in this snapshot

# Column names that must never appear in a generated source file --
# Ground-Truth-only / analytical-conclusion fields, per the Source &
# Delivery Model checkpoint's leakage prohibition.
PROHIBITED_COLUMN_NAMES = {
    "unmet_demand",
    "stockout",
    "potential_demand",
    "promotion_id",
    "promotion_phase",
    "physical_inventory",
    "opening_inventory",
    "closing_inventory",
    "store_scale_class",
    "sell_eligible",
    "available",
    "store_id",  # canonical Phase 1 identifier -- must never appear as-is
    "sku_id",  # canonical Phase 1 identifier -- must never appear as-is
}


# ---------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------


@pytest.fixture(scope="module")
def snapshot():
    return read_phase1_snapshot(SNAPSHOT_DIR)


@pytest.fixture(scope="module")
def profile_a(snapshot):
    return build_profile_a(snapshot.retailer_id, snapshot.store_ids, snapshot.sku_ids)


@pytest.fixture(scope="module")
def profile_b(snapshot):
    return build_profile_b(
        snapshot.retailer_id, snapshot.store_ids, snapshot.sku_ids, snapshot.run_start_date, seed=42
    )


def _all_dates(snapshot) -> list:
    num_days = (snapshot.run_end_date - snapshot.run_start_date).days + 1
    return [snapshot.run_start_date + timedelta(days=offset) for offset in range(num_days)]


def _weekly_periods(snapshot) -> list:
    num_days = (snapshot.run_end_date - snapshot.run_start_date).days + 1
    return weekly_delivery_periods(snapshot.run_start_date, num_days)


def _hash_tree(root: Path) -> dict[str, str]:
    return {
        str(path.relative_to(root)): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def _reverse(mapping: dict[str, str]) -> dict[str, str]:
    return {v: k for k, v in mapping.items()}


def _read_csv_header(path: Path) -> list[str]:
    with path.open(encoding="utf-8") as handle:
        return next(csv.reader(handle))


def _read_xlsx_header(path: Path) -> list[str]:
    workbook = load_workbook(path)
    return [cell.value for cell in next(workbook.active.iter_rows(min_row=1, max_row=1))]


# ---------------------------------------------------------------------
# 1 / 24: reproducibility -- same profile + seed -> identical, repeatable bytes
# ---------------------------------------------------------------------


def test_same_profile_and_seed_produce_byte_identical_artifacts(snapshot, tmp_path):
    profile_b_1 = build_profile_b(
        snapshot.retailer_id, snapshot.store_ids, snapshot.sku_ids, snapshot.run_start_date, seed=7
    )
    profile_b_2 = build_profile_b(
        snapshot.retailer_id, snapshot.store_ids, snapshot.sku_ids, snapshot.run_start_date, seed=7
    )

    root_1 = write_generated_feeds(
        tmp_path / "run1", "profile_b", generate_profile_feeds(snapshot, profile_b_1)
    )
    root_2 = write_generated_feeds(
        tmp_path / "run2", "profile_b", generate_profile_feeds(snapshot, profile_b_2)
    )

    hashes_1 = _hash_tree(root_1)
    hashes_2 = _hash_tree(root_2)
    assert hashes_1 == hashes_2
    assert len(hashes_1) > 0


def test_repeated_execution_is_reproducible(snapshot, tmp_path):
    """Same as above, phrased as its own invariant (design brief #24):
    running generation a third time still matches."""
    profile = build_profile_a(snapshot.retailer_id, snapshot.store_ids, snapshot.sku_ids)
    roots = [
        write_generated_feeds(
            tmp_path / f"run{i}", "profile_a", generate_profile_feeds(snapshot, profile)
        )
        for i in range(3)
    ]
    hashes = [_hash_tree(root) for root in roots]
    assert hashes[0] == hashes[1] == hashes[2]


# ---------------------------------------------------------------------
# 2: different seeds only alter stochastic fault realization
# ---------------------------------------------------------------------


def test_different_seed_changes_only_the_stochastic_excess_latency_day(snapshot):
    profile_1 = build_profile_b(
        snapshot.retailer_id, snapshot.store_ids, snapshot.sku_ids, snapshot.run_start_date, seed=1
    )
    profile_2 = build_profile_b(
        snapshot.retailer_id, snapshot.store_ids, snapshot.sku_ids, snapshot.run_start_date, seed=2
    )
    faults_1 = profile_1.feeds["POS"].fault_profile
    faults_2 = profile_2.feeds["POS"].fault_profile

    # structural characteristics and every other fault category unchanged
    assert (
        profile_1.feeds["POS"].baseline_latency_days == profile_2.feeds["POS"].baseline_latency_days
    )
    assert profile_1.feeds["POS"].format == profile_2.feeds["POS"].format
    assert faults_1.missing_delivery_periods == faults_2.missing_delivery_periods
    assert faults_1.duplicate_delivery_periods == faults_2.duplicate_delivery_periods
    assert faults_1.store_migration == faults_2.store_migration

    # the one stochastic element stays within its configured bounds
    eligible_dates = {
        snapshot.run_start_date + timedelta(days=o) for o in EXCESS_LATENCY_ELIGIBLE_OFFSETS
    }
    assert set(faults_1.excess_latency_days) <= eligible_dates
    assert set(faults_2.excess_latency_days) <= eligible_dates
    assert set(faults_1.excess_latency_days.values()) == {EXCESS_LATENCY_EXTRA_DAYS}

    # and it does genuinely vary across seeds (not vacuously constant)
    days_seen = set()
    for seed in range(20):
        profile = build_profile_b(
            snapshot.retailer_id,
            snapshot.store_ids,
            snapshot.sku_ids,
            snapshot.run_start_date,
            seed=seed,
        )
        days_seen.update(profile.feeds["POS"].fault_profile.excess_latency_days)
    assert days_seen <= eligible_dates
    assert len(days_seen) > 1


# ---------------------------------------------------------------------
# 3: Phase 1 snapshot is never modified by Phase 2 generation
# ---------------------------------------------------------------------


def test_phase1_snapshot_unchanged_after_both_profiles_run(tmp_path):
    before = _hash_tree(SNAPSHOT_DIR)
    run_vertical_slice_1(SNAPSHOT_DIR, tmp_path, phase2_seed=42)
    after = _hash_tree(SNAPSHOT_DIR)
    assert before == after


def test_snapshot_reader_never_requests_hidden_ground_truth_columns():
    """Structural guarantee, not just a runtime check: the Parquet
    columns Phase 2 ever asks for are enumerated explicitly and never
    include Document 8's evaluation-only fields.
    """
    assert "unmet_demand" not in _DAILY_STATE_COLUMNS
    assert "stockout" not in _DAILY_STATE_COLUMNS
    assert "potential_demand" not in _DAILY_STATE_COLUMNS
    assert "promotion_id" not in _DAILY_STATE_COLUMNS


# ---------------------------------------------------------------------
# 4 / 5 / 6: POS row-existence gate and provenance (Document 7 boundary)
# ---------------------------------------------------------------------


def test_never_sell_eligible_pair_produces_no_pos_row_ever(snapshot, profile_a):
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts, _all_dates(snapshot), profile_a.identifier_scheme, 1, None
    )
    store_id, sku_id = NEVER_CARRIED_PAIR
    assert not any(
        f.store_id == store_id and f.sku_id == sku_id and f.sell_eligible
        for f in snapshot.pos_facts
    )
    store_code = profile_a.identifier_scheme.store_mapping[store_id]
    item_code = profile_a.identifier_scheme.product_mapping[sku_id]
    for artifact in artifacts:
        assert not any(
            r.store_code == store_code and r.item_code == item_code for r in artifact.rows
        )


def test_legitimate_zero_sales_day_produces_a_real_zero_row(snapshot, profile_a):
    zero_sales_fact = next(f for f in snapshot.pos_facts if f.sell_eligible and f.sales == 0)
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts, _all_dates(snapshot), profile_a.identifier_scheme, 1, None
    )
    artifact = next(a for a in artifacts if a.delivery_period == zero_sales_fact.date)
    store_code = profile_a.identifier_scheme.store_mapping[zero_sales_fact.store_id]
    item_code = profile_a.identifier_scheme.product_mapping[zero_sales_fact.sku_id]
    matching = [r for r in artifact.rows if r.store_code == store_code and r.item_code == item_code]
    assert len(matching) == 1
    assert matching[0].sales_qty == 0


def test_every_pos_row_traces_to_a_phase1_sales_fact_never_invented(snapshot, profile_a):
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts, _all_dates(snapshot), profile_a.identifier_scheme, 1, None
    )
    reverse_store = _reverse(profile_a.identifier_scheme.store_mapping)
    reverse_sku = _reverse(profile_a.identifier_scheme.product_mapping)
    truth = {(f.store_id, f.sku_id, f.date): f.sales for f in snapshot.pos_facts if f.sell_eligible}

    observed = {}
    for artifact in artifacts:
        for row in artifact.rows:
            key = (reverse_store[row.store_code], reverse_sku[row.item_code], row.business_date)
            observed[key] = row.sales_qty

    assert (
        observed == truth
    )  # every reported figure equals a real Phase 1 fact, nothing more, nothing less


# ---------------------------------------------------------------------
# 7 / 8: identifier stability and bijectivity
# ---------------------------------------------------------------------


def test_profile_a_identifiers_deterministic_stable_and_bijective(snapshot):
    profile_1 = build_profile_a(snapshot.retailer_id, snapshot.store_ids, snapshot.sku_ids)
    profile_2 = build_profile_a(snapshot.retailer_id, snapshot.store_ids, snapshot.sku_ids)
    assert profile_1.identifier_scheme.store_mapping == profile_2.identifier_scheme.store_mapping
    assert (
        profile_1.identifier_scheme.product_mapping == profile_2.identifier_scheme.product_mapping
    )

    store_values = list(profile_1.identifier_scheme.store_mapping.values())
    sku_values = list(profile_1.identifier_scheme.product_mapping.values())
    assert len(store_values) == len(set(store_values))
    assert len(sku_values) == len(set(sku_values))
    # retailer-local, never NovaFoods canonical
    assert set(profile_1.identifier_scheme.store_mapping.values()).isdisjoint(snapshot.store_ids)


def test_profile_b_product_identifiers_untouched_by_store_migration(snapshot, profile_b):
    migration = profile_b.feeds["POS"].fault_profile.store_migration
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts,
        _all_dates(snapshot),
        profile_b.identifier_scheme,
        3,
        profile_b.feeds["POS"].fault_profile,
    )
    stable_sku_values = set(profile_b.identifier_scheme.product_mapping.values())
    for artifact in artifacts:
        for row in artifact.rows:
            if row.store_code in (migration.old_identifier, migration.new_identifier):
                assert row.item_code in stable_sku_values


# ---------------------------------------------------------------------
# 9 / 10 / 11 / 12: store identifier migration
# ---------------------------------------------------------------------


def test_resolve_store_identifier_old_then_gap_then_new_never_overlapping(profile_b):
    migration = profile_b.feeds["POS"].fault_profile.store_migration
    for offset in range(-3, 10):
        day = migration.gap_start + timedelta(days=offset)
        result = resolve_store_identifier(
            migration.store_id, day, profile_b.identifier_scheme.store_mapping, migration
        )
        if migration.gap_start <= day <= migration.gap_end:
            assert result is None
        elif day < migration.gap_start:
            assert result == migration.old_identifier
        else:
            assert result == migration.new_identifier


def test_store_migration_gap_excludes_the_store_from_every_artifact(snapshot, profile_b):
    faults = profile_b.feeds["POS"].fault_profile
    migration = faults.store_migration
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts, _all_dates(snapshot), profile_b.identifier_scheme, 3, faults
    )

    before_ids, gap_ids, after_ids = set(), set(), set()
    for artifact in artifacts:
        ids_present = {r.store_code for r in artifact.rows}
        if artifact.delivery_period < migration.gap_start:
            before_ids |= ids_present
            assert migration.new_identifier not in ids_present
        elif migration.gap_start <= artifact.delivery_period <= migration.gap_end:
            gap_ids |= ids_present
            assert migration.old_identifier not in ids_present
            assert migration.new_identifier not in ids_present
        else:
            after_ids |= ids_present
            assert migration.old_identifier not in ids_present

    assert migration.old_identifier in before_ids
    assert migration.new_identifier in after_ids


def test_migration_affects_every_sku_at_the_store_consistently(snapshot, profile_b):
    faults = profile_b.feeds["POS"].fault_profile
    migration = faults.store_migration
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts, _all_dates(snapshot), profile_b.identifier_scheme, 3, faults
    )
    reverse_sku = _reverse(profile_b.identifier_scheme.product_mapping)

    after_gap_facts = [
        f
        for f in snapshot.pos_facts
        if f.store_id == migration.store_id and f.sell_eligible and f.date > migration.gap_end
    ]
    assert after_gap_facts  # sanity: there is real data to check this against
    sample_date = after_gap_facts[0].date
    expected_skus = {f.sku_id for f in after_gap_facts if f.date == sample_date}

    artifact = next(a for a in artifacts if a.delivery_period == sample_date)
    store_rows = [
        r
        for r in artifact.rows
        if r.store_code in (migration.old_identifier, migration.new_identifier)
    ]
    assert all(r.store_code == migration.new_identifier for r in store_rows)
    assert {reverse_sku[r.item_code] for r in store_rows} == expected_skus


def test_migration_is_the_identical_fault_instance_across_pos_and_assortment(profile_b):
    pos_migration = profile_b.feeds["POS"].fault_profile.store_migration
    assortment_migration = profile_b.feeds["ASSORTMENT"].fault_profile.store_migration
    assert pos_migration is assortment_migration


def test_migration_gap_excludes_store_from_assortment_snapshot_too(snapshot, profile_b):
    faults = profile_b.feeds["ASSORTMENT"].fault_profile
    migration = faults.store_migration
    artifacts = generate_assortment_artifacts(
        snapshot.assortment_facts,
        _weekly_periods(snapshot),
        snapshot.run_start_date,
        profile_b.identifier_scheme,
        7,
        faults,
    )
    gap_period = next(
        p for p in _weekly_periods(snapshot) if migration.gap_start <= p <= migration.gap_end
    )
    artifact = next(a for a in artifacts if a.delivery_period == gap_period)
    ids_present = {r.store_code for r in artifact.rows}
    assert migration.old_identifier not in ids_present
    assert migration.new_identifier not in ids_present


# ---------------------------------------------------------------------
# 13 / 14 / 15: POS operational faults (missing / duplicate / latency)
# ---------------------------------------------------------------------


def test_missing_pos_delivery_creates_no_landing_directory(snapshot, profile_b, tmp_path):
    generated = generate_profile_feeds(snapshot, profile_b)
    root = write_generated_feeds(tmp_path, "profile_b", generated)
    for offset in MISSING_POS_DELIVERY_OFFSETS:
        missing_day = snapshot.run_start_date + timedelta(days=offset)
        assert not (root / "POS" / missing_day.isoformat()).exists()


def test_duplicate_pos_delivery_is_byte_identical_on_disk(snapshot, profile_b, tmp_path):
    generated = generate_profile_feeds(snapshot, profile_b)
    root = write_generated_feeds(tmp_path, "profile_b", generated)
    dup_day = snapshot.run_start_date + timedelta(days=DUPLICATE_POS_DELIVERY_OFFSET)
    period_dir = root / "POS" / dup_day.isoformat()
    original = (period_dir / "pos.csv").read_bytes()
    duplicate = (period_dir / "pos.duplicate.csv").read_bytes()
    assert original == duplicate
    assert len(original) > 0


def test_excess_latency_applies_only_to_its_configured_day_and_amount(snapshot, profile_b):
    faults = profile_b.feeds["POS"].fault_profile
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts, _all_dates(snapshot), profile_b.identifier_scheme, 3, faults
    )
    eligible_dates = {
        snapshot.run_start_date + timedelta(days=o) for o in EXCESS_LATENCY_ELIGIBLE_OFFSETS
    }
    assert set(faults.excess_latency_days) <= eligible_dates

    for artifact in artifacts:
        expected_extra = faults.excess_latency_days.get(artifact.delivery_period, 0)
        assert artifact.delivered_on == artifact.delivery_period + timedelta(
            days=3 + expected_extra
        )


# ---------------------------------------------------------------------
# 16 / 17: Assortment feed -- staleness and missing delivery
# ---------------------------------------------------------------------


def test_delayed_assortment_visibility_traces_to_the_real_phase1_assortment_change(
    snapshot, profile_a, profile_b
):
    periods = _weekly_periods(snapshot)
    a_artifacts = generate_assortment_artifacts(
        snapshot.assortment_facts,
        periods,
        snapshot.run_start_date,
        profile_a.identifier_scheme,
        2,
        None,
    )
    b_faults = profile_b.feeds["ASSORTMENT"].fault_profile
    b_artifacts = generate_assortment_artifacts(
        snapshot.assortment_facts,
        periods,
        snapshot.run_start_date,
        profile_b.identifier_scheme,
        7,
        b_faults,
    )

    store_id, sku_id = TEMPORAL_GAP_PAIR
    pair_key_a = (
        profile_a.identifier_scheme.store_mapping[store_id],
        profile_a.identifier_scheme.product_mapping[sku_id],
    )
    pair_key_b = (
        profile_b.identifier_scheme.store_mapping[store_id],
        profile_b.identifier_scheme.product_mapping[sku_id],
    )

    def carried(artifacts, pair_key, period):
        artifact = next(a for a in artifacts if a.delivery_period == period)
        return pair_key in {(r.store_code, r.item_code) for r in artifact.rows}

    period_35 = snapshot.run_start_date + timedelta(days=35)
    assert carried(a_artifacts, pair_key_a, period_35) is False  # A already reflects the real drop
    assert carried(b_artifacts, pair_key_b, period_35) is True  # B is still lagging behind it

    b_artifact_35 = next(a for a in b_artifacts if a.delivery_period == period_35)
    assert b_artifact_35.source_as_of_date == period_35 - timedelta(
        days=STALE_ASSORTMENT_VISIBILITY_EXTRA_DAYS
    )
    assert any(
        f.store_id == store_id
        and f.sku_id == sku_id
        and f.date == b_artifact_35.source_as_of_date
        and f.physical_assortment
        for f in snapshot.assortment_facts
    )  # B's reported state is exactly Phase 1's truth at the lagged source date -- not invented


def test_missing_assortment_delivery_creates_no_landing_directory(snapshot, profile_b, tmp_path):
    generated = generate_profile_feeds(snapshot, profile_b)
    root = write_generated_feeds(tmp_path, "profile_b", generated)
    missing_period = snapshot.run_start_date + timedelta(days=MISSING_ASSORTMENT_DELIVERY_OFFSET)
    assert not (root / "ASSORTMENT" / missing_period.isoformat()).exists()


def test_every_assortment_row_traces_to_phase1_physical_assortment_never_invented(
    snapshot, profile_a
):
    periods = _weekly_periods(snapshot)
    artifacts = generate_assortment_artifacts(
        snapshot.assortment_facts,
        periods,
        snapshot.run_start_date,
        profile_a.identifier_scheme,
        2,
        None,
    )
    reverse_store = _reverse(profile_a.identifier_scheme.store_mapping)
    reverse_sku = _reverse(profile_a.identifier_scheme.product_mapping)

    truth_by_date = defaultdict(set)
    for fact in snapshot.assortment_facts:
        if fact.physical_assortment:
            truth_by_date[fact.date].add((fact.store_id, fact.sku_id))

    for artifact in artifacts:
        reported = {(reverse_store[r.store_code], reverse_sku[r.item_code]) for r in artifact.rows}
        assert reported == truth_by_date.get(artifact.source_as_of_date, set())


# ---------------------------------------------------------------------
# 18: weekly POS aggregation -- not exercised in this slice, recorded explicitly
# ---------------------------------------------------------------------


def test_pos_stays_at_store_sku_day_grain_for_both_profiles_this_slice(profile_a, profile_b):
    """Section 7: both Profile A and Profile B report POS at store x
    SKU x day. Neither exercises weekly POS aggregation in this slice,
    so there is no daily-to-weekly rollup to test here -- recorded
    explicitly (design brief invariant #5/#18) rather than left silently
    unverified.
    """
    assert profile_a.feeds["POS"].aggregation_level == "store_sku_day"
    assert profile_b.feeds["POS"].aggregation_level == "store_sku_day"


# ---------------------------------------------------------------------
# 19: structural feed absence
# ---------------------------------------------------------------------


def test_structural_feed_absence_produces_no_artifacts_and_no_directory(snapshot, tmp_path):
    pos_only_profile = RetailerFeedProfile(
        retailer_id=snapshot.retailer_id,
        identifier_scheme=build_identifier_scheme(snapshot.store_ids, snapshot.sku_ids),
        feeds={
            "POS": FeedProfile(
                source_family=FEED_SOURCE_FAMILY["POS"],
                delivery_mechanism=FEED_DELIVERY_MECHANISM["POS"],
                format="csv",
                frequency="daily",
                aggregation_level="store_sku_day",
                baseline_latency_days=1,
                column_casing=COLUMN_CASING_LOWER_SNAKE_CASE,
                fault_profile=None,
            )
        },
    )
    generated = generate_profile_feeds(snapshot, pos_only_profile)
    assert generated.assortment_artifacts == []

    root = write_generated_feeds(tmp_path, "pos_only", generated)
    assert (root / "POS").exists()
    assert not (root / "ASSORTMENT").exists()


# ---------------------------------------------------------------------
# 20: every realized fault lies exactly inside its configured scope
# ---------------------------------------------------------------------


def test_pos_faults_are_scoped_exactly_as_configured_no_more_no_less(snapshot, profile_b):
    faults = profile_b.feeds["POS"].fault_profile
    all_dates = _all_dates(snapshot)
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts, all_dates, profile_b.identifier_scheme, 3, faults
    )

    delivered_periods = {a.delivery_period for a in artifacts if not a.is_duplicate}
    expected_missing = {
        snapshot.run_start_date + timedelta(days=o) for o in MISSING_POS_DELIVERY_OFFSETS
    }
    assert set(all_dates) - delivered_periods == expected_missing

    duplicate_periods = {a.delivery_period for a in artifacts if a.is_duplicate}
    assert duplicate_periods == {
        snapshot.run_start_date + timedelta(days=DUPLICATE_POS_DELIVERY_OFFSET)
    }


# ---------------------------------------------------------------------
# 21: Profile A matches Phase 1 through its declared mapping/cadence/latency
# ---------------------------------------------------------------------


def test_profile_a_reproduces_phase1_pos_exactly_through_its_declared_transform(
    snapshot, profile_a
):
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts, _all_dates(snapshot), profile_a.identifier_scheme, 1, None
    )
    reverse_store = _reverse(profile_a.identifier_scheme.store_mapping)
    reverse_sku = _reverse(profile_a.identifier_scheme.product_mapping)

    observed = {
        (
            reverse_store[row.store_code],
            reverse_sku[row.item_code],
            row.business_date,
        ): row.sales_qty
        for artifact in artifacts
        for row in artifact.rows
    }
    truth = {(f.store_id, f.sku_id, f.date): f.sales for f in snapshot.pos_facts if f.sell_eligible}
    assert observed == truth
    assert all(a.delivered_on == a.delivery_period + timedelta(days=1) for a in artifacts)
    assert len(artifacts) == len(_all_dates(snapshot))  # no missing deliveries in the clean profile


# ---------------------------------------------------------------------
# 22: Profile B differs from Profile A only through explicitly configured mechanisms
# ---------------------------------------------------------------------


def test_profile_b_diverges_from_profile_a_only_via_named_faults(snapshot, profile_a, profile_b):
    all_dates = _all_dates(snapshot)
    a_artifacts = {
        a.delivery_period: a
        for a in generate_pos_artifacts(
            snapshot.pos_facts, all_dates, profile_a.identifier_scheme, 1, None
        )
    }
    b_faults = profile_b.feeds["POS"].fault_profile
    b_artifacts = {
        a.delivery_period: a
        for a in generate_pos_artifacts(
            snapshot.pos_facts, all_dates, profile_b.identifier_scheme, 3, b_faults
        )
        if not a.is_duplicate
    }

    reverse_store_a = _reverse(profile_a.identifier_scheme.store_mapping)
    reverse_sku_a = _reverse(profile_a.identifier_scheme.product_mapping)
    reverse_store_b = _reverse(profile_b.identifier_scheme.store_mapping)
    reverse_sku_b = _reverse(profile_b.identifier_scheme.product_mapping)
    migration = b_faults.store_migration

    for period, a_artifact in a_artifacts.items():
        if period in b_faults.missing_delivery_periods:
            assert period not in b_artifacts  # the one named "missing delivery" mechanism
            continue

        b_artifact = b_artifacts[period]
        a_canonical = {
            (reverse_store_a[r.store_code], reverse_sku_a[r.item_code]): r.sales_qty
            for r in a_artifact.rows
        }
        b_canonical = {}
        for row in b_artifact.rows:
            store_id = (
                migration.store_id
                if row.store_code == migration.new_identifier
                else reverse_store_b[row.store_code]
            )
            b_canonical[(store_id, reverse_sku_b[row.item_code])] = row.sales_qty

        if migration.gap_start <= period <= migration.gap_end:
            expected = {k: v for k, v in a_canonical.items() if k[0] != migration.store_id}
        else:
            expected = a_canonical
        assert b_canonical == expected  # every remaining divergence is the migration, nothing else


# ---------------------------------------------------------------------
# 23: independently generated Profile A / Profile B artifact trees
# ---------------------------------------------------------------------


def test_profile_a_and_profile_b_produce_independent_landing_trees(snapshot, tmp_path):
    results = run_vertical_slice_1(SNAPSHOT_DIR, tmp_path, phase2_seed=42)
    root_a, root_b = results["profile_a"], results["profile_b"]
    assert root_a != root_b

    missing_day = snapshot.run_start_date + timedelta(days=MISSING_POS_DELIVERY_OFFSETS[0])
    assert (root_a / "POS" / missing_day.isoformat()).exists()
    assert not (root_b / "POS" / missing_day.isoformat()).exists()


# ---------------------------------------------------------------------
# Written-artifact round trips (real files on disk, not just in-memory objects)
# ---------------------------------------------------------------------


def test_written_pos_csv_round_trips_to_the_same_row_count(snapshot, profile_a, tmp_path):
    generated = generate_profile_feeds(snapshot, profile_a)
    root = write_generated_feeds(tmp_path, "profile_a", generated)
    sample_artifact = generated.pos_artifacts[10]
    period_dir = root / "POS" / sample_artifact.delivery_period.isoformat()
    with (period_dir / "pos.csv").open(encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == len(sample_artifact.rows)
    assert list(rows[0].keys()) == list(POS_COLUMN_ORDER)


def test_written_assortment_xlsx_round_trips_to_the_same_row_count(snapshot, profile_b, tmp_path):
    generated = generate_profile_feeds(snapshot, profile_b)
    root = write_generated_feeds(tmp_path, "profile_b", generated)
    sample_artifact = generated.assortment_artifacts[0]
    period_dir = root / "ASSORTMENT" / sample_artifact.delivery_period.isoformat()
    workbook = load_workbook(period_dir / "assortment.xlsx")
    sheet = workbook.active
    data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == len(sample_artifact.rows)


# ---------------------------------------------------------------------
# Source & Delivery Model checkpoint: source_family / delivery_mechanism
# ---------------------------------------------------------------------


def test_source_family_is_fixed_per_feed_type_regardless_of_profile(profile_a, profile_b):
    for profile in (profile_a, profile_b):
        assert profile.feeds["POS"].source_family == "transactional"
        assert profile.feeds["ASSORTMENT"].source_family == "business_document"


def test_delivery_mechanism_is_fixed_per_feed_type_regardless_of_profile(profile_a, profile_b):
    for profile in (profile_a, profile_b):
        assert profile.feeds["POS"].delivery_mechanism == "cloud_object_landing"
        assert profile.feeds["ASSORTMENT"].delivery_mechanism == "document_repository"


# ---------------------------------------------------------------------
# Source & Delivery Model checkpoint: approved operational-baggage fields present
# ---------------------------------------------------------------------


def test_pos_schema_contains_the_approved_fields_in_the_approved_order():
    assert POS_COLUMN_ORDER == (
        "store_code",
        "item_code",
        "business_date",
        "sales_qty",
        "uom",
        "currency_code",
        "exported_at",
    )


def test_assortment_schema_contains_the_approved_fields_in_the_approved_order():
    assert ASSORTMENT_COLUMN_ORDER == (
        "store_code",
        "item_code",
        "retailer_category_code",
        "item_description",
        "record_status_code",
        "exported_at",
    )


def test_pos_operational_baggage_values_are_constant_and_do_not_encode_ground_truth(
    snapshot, profile_a
):
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts, _all_dates(snapshot), profile_a.identifier_scheme, 1, None
    )
    uoms = {row.uom for artifact in artifacts for row in artifact.rows}
    currencies = {row.currency_code for artifact in artifacts for row in artifact.rows}
    assert uoms == {"EA"}
    assert currencies == {"USD"}


def test_assortment_operational_baggage_values_are_constant_and_do_not_encode_ground_truth(
    snapshot, profile_a
):
    artifacts = generate_assortment_artifacts(
        snapshot.assortment_facts,
        _weekly_periods(snapshot),
        snapshot.run_start_date,
        profile_a.identifier_scheme,
        2,
        None,
    )
    statuses = {row.record_status_code for artifact in artifacts for row in artifact.rows}
    assert statuses == {"A"}  # constant legacy field -- never varies with physical_assortment
    # descriptions are a pure function of item_code, never of any hidden attribute
    for artifact in artifacts:
        for row in artifact.rows:
            assert row.item_description == f"NovaFoods Product {row.item_code}"


def test_exported_at_is_retailer_side_metadata_computed_from_delivery_period_only(snapshot):
    profile_a = build_profile_a(snapshot.retailer_id, snapshot.store_ids, snapshot.sku_ids)
    artifacts = generate_pos_artifacts(
        snapshot.pos_facts, _all_dates(snapshot), profile_a.identifier_scheme, 1, None
    )
    for artifact in artifacts:
        if not artifact.rows:
            continue
        expected = datetime.combine(artifact.delivery_period + timedelta(days=1), time(2, 13))
        assert all(row.exported_at == expected for row in artifact.rows)
        # exported_at (artifact/retailer-side metadata) is a wholly different
        # mechanism from delivered_on (delivery-event/NovaFoods-side
        # metadata) -- computed independently, never derived from each other
        assert isinstance(artifact.delivered_on, type(artifact.delivery_period))
        assert isinstance(artifact.rows[0].exported_at, datetime)


def test_prohibited_ground_truth_column_names_never_appear_in_generated_schemas():
    assert PROHIBITED_COLUMN_NAMES.isdisjoint(POS_COLUMN_ORDER)
    assert PROHIBITED_COLUMN_NAMES.isdisjoint(ASSORTMENT_COLUMN_ORDER)


# ---------------------------------------------------------------------
# Source & Delivery Model checkpoint: Profile A/B column casing (Section 7)
# ---------------------------------------------------------------------


def test_profile_a_pos_header_is_lower_snake_case(snapshot, profile_a, tmp_path):
    generated = generate_profile_feeds(snapshot, profile_a)
    root = write_generated_feeds(tmp_path, "profile_a", generated)
    sample = generated.pos_artifacts[10]
    header = _read_csv_header(root / "POS" / sample.delivery_period.isoformat() / "pos.csv")
    assert header == list(POS_COLUMN_ORDER)


def test_profile_b_pos_header_is_upper_snake_case(snapshot, profile_b, tmp_path):
    generated = generate_profile_feeds(snapshot, profile_b)
    root = write_generated_feeds(tmp_path, "profile_b", generated)
    non_missing = [a for a in generated.pos_artifacts if not a.is_duplicate]
    sample = non_missing[10]
    header = _read_csv_header(root / "POS" / sample.delivery_period.isoformat() / "pos.csv")
    assert header == [c.upper() for c in POS_COLUMN_ORDER]


def test_profile_a_assortment_header_is_lower_snake_case(snapshot, profile_a, tmp_path):
    generated = generate_profile_feeds(snapshot, profile_a)
    root = write_generated_feeds(tmp_path, "profile_a", generated)
    sample = generated.assortment_artifacts[0]
    header = _read_csv_header(
        root / "ASSORTMENT" / sample.delivery_period.isoformat() / "assortment.csv"
    )
    assert header == list(ASSORTMENT_COLUMN_ORDER)


def test_profile_b_assortment_header_is_upper_snake_case(snapshot, profile_b, tmp_path):
    generated = generate_profile_feeds(snapshot, profile_b)
    root = write_generated_feeds(tmp_path, "profile_b", generated)
    sample = generated.assortment_artifacts[0]
    header = _read_xlsx_header(
        root / "ASSORTMENT" / sample.delivery_period.isoformat() / "assortment.xlsx"
    )
    assert header == [c.upper() for c in ASSORTMENT_COLUMN_ORDER]


def test_profile_a_and_b_pos_schema_identical_shape_and_order_no_abbreviation_divergence(
    snapshot, profile_a, profile_b, tmp_path
):
    generated_a = generate_profile_feeds(snapshot, profile_a)
    generated_b = generate_profile_feeds(snapshot, profile_b)
    root_a = write_generated_feeds(tmp_path / "a", "profile_a", generated_a)
    root_b = write_generated_feeds(tmp_path / "b", "profile_b", generated_b)

    header_a = _read_csv_header(
        root_a / "POS" / generated_a.pos_artifacts[10].delivery_period.isoformat() / "pos.csv"
    )
    non_missing_b = [a for a in generated_b.pos_artifacts if not a.is_duplicate]
    header_b = _read_csv_header(
        root_b / "POS" / non_missing_b[10].delivery_period.isoformat() / "pos.csv"
    )

    assert len(header_a) == len(header_b)
    # the ONLY approved difference is casing -- uppercasing A's header must
    # exactly equal B's header, proving no field was also abbreviated/renamed
    assert [c.upper() for c in header_a] == header_b


def test_profile_a_and_b_assortment_schema_identical_shape_and_order_no_abbreviation_divergence(
    snapshot, profile_a, profile_b, tmp_path
):
    generated_a = generate_profile_feeds(snapshot, profile_a)
    generated_b = generate_profile_feeds(snapshot, profile_b)
    root_a = write_generated_feeds(tmp_path / "a", "profile_a", generated_a)
    root_b = write_generated_feeds(tmp_path / "b", "profile_b", generated_b)

    header_a = _read_csv_header(
        root_a
        / "ASSORTMENT"
        / generated_a.assortment_artifacts[0].delivery_period.isoformat()
        / "assortment.csv"
    )
    header_b = _read_xlsx_header(
        root_b
        / "ASSORTMENT"
        / generated_b.assortment_artifacts[0].delivery_period.isoformat()
        / "assortment.xlsx"
    )

    assert len(header_a) == len(header_b)
    assert [c.upper() for c in header_a] == header_b
